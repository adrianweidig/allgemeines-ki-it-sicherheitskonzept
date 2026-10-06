import hashlib
import shutil
import subprocess
import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'validierung'))
import erzeuge_dokumente as ed


class Nutzerbelehrung(unittest.TestCase):
    def test_referenzdokumente_ohne_persoenliche_metadaten(self):
        for pfad in (ed.ZIEL, ed.BELEHRUNG_ZIEL):
            with self.subTest(dokument=pfad.name):
                eigenschaften = Document(pfad).core_properties
                pdf = PdfReader(pfad.with_suffix('.pdf'))
                self.assertFalse(eigenschaften.author)
                self.assertFalse(eigenschaften.last_modified_by)
                self.assertFalse(pdf.metadata.author)
                self.assertTrue(eigenschaften.title)
                self.assertEqual(eigenschaften.title, pdf.metadata.title)

    def test_vorlage_felder_und_gespeicherte_werte(self):
        reader = PdfReader(ed.BELEHRUNG_ZIEL.with_suffix('.pdf'))
        self.assertIn(len(reader.pages), (1, 2))
        self.assertEqual(ed.BELEHRUNG_TITEL, reader.metadata.title)
        self.assertFalse(reader.metadata.author)
        fields = reader.get_fields()
        self.assertEqual({n for n, _, _ in ed.BELEHRUNG_FELDER}, set(fields))
        self.assertEqual('/Sig', fields['unterschrift']['/FT'])
        self.assertNotIn('/V', fields['unterschrift'])
        widgets = [a.get_object() for p in reader.pages for a in p.get('/Annots', [])]
        self.assertEqual(4, len(widgets))
        for i, widget in enumerate(widgets):
            self.assertIn(widget.indirect_reference, reader.trailer['/Root']['/AcroForm']['/Fields'])
            self.assertTrue(widget['/AP']['/N'].get_data())
            self.assertNotIn('/A', widget)
            self.assertNotIn('/AA', widget)
            ax0, ay0, ax1, ay1 = widget['/Rect']
            for other in widgets[i+1:]:
                bx0, by0, bx1, by1 = other['/Rect']
                self.assertTrue(ax1 <= bx0 or bx1 <= ax0 or ay1 <= by0 or by1 <= ay0)
        signature = next(w for w in widgets if w['/FT'] == '/Sig')
        self.assertEqual('/All', signature['/Lock']['/Action'])
        doc = Document(ed.BELEHRUNG_ZIEL)
        text = '\n'.join(p.text for p in doc.paragraphs)
        for required in ('Cloud-Nutzung', 'regulären Nutzersitzung', 'Einstufung eigener Informationen'):
            self.assertIn(required, text)
        for forbidden in ('vor der ersten', 'Digitale Unterschrift', 'Autor:', ';'):
            self.assertNotIn(forbidden, text)
        self.assertFalse(doc.core_properties.author)
        (ROOT / '.arbeitsdaten').mkdir(exist_ok=True)
        with TemporaryDirectory(dir=ROOT / '.arbeitsdaten') as temp:
            path = Path(temp) / 'gefüllt.pdf'
            values = {n: 'Prüfperson ÄÖÜ äöü ß – technischer Test' for n in fields if n != 'unterschrift'}
            writer = PdfWriter(clone_from=reader)
            writer.update_page_form_field_values(None, values, auto_regenerate=False)
            writer.write(path)
            saved = PdfReader(path)
            for n, value in values.items():
                self.assertEqual(value, saved.get_fields()[n]['/V'])
            for page in saved.pages:
                for ref in page.get('/Annots', []):
                    widget = ref.get_object()
                    if widget['/FT'] == '/Tx':
                        self.assertEqual(values[widget['/T']], widget['/V'])
                        self.assertIn(b'Tj', widget['/AP']['/N'].get_data())

    def test_formularaufbau_überschreibt_keine_bestehenden_felder(self):
        (ROOT / '.arbeitsdaten').mkdir(exist_ok=True)
        with TemporaryDirectory(dir=ROOT / '.arbeitsdaten') as temp:
            path = Path(temp) / 'export.pdf'
            writer = PdfWriter(clone_from=ed.BELEHRUNG_ZIEL.with_suffix('.pdf'))
            writer.remove_annotations(subtypes='/Widget')
            writer.root_object.pop(NameObject('/AcroForm'))
            writer.write(path)
            before = hashlib.sha256(path.read_bytes()).digest()
            with patch.object(ed, 'BELEHRUNG_FELDER', (('fehlt', 'Nicht vorhandene Beschriftung', 24),)):
                with self.assertRaisesRegex(ValueError, 'nicht eindeutig'):
                    ed.ergänze_belehrungsformular(path)
            self.assertEqual(before, hashlib.sha256(path.read_bytes()).digest())
            ed.ergänze_belehrungsformular(path)
            before = hashlib.sha256(path.read_bytes()).digest()
            with self.assertRaisesRegex(ValueError, 'ohne Formularfelder'):
                ed.ergänze_belehrungsformular(path)
            self.assertEqual(before, hashlib.sha256(path.read_bytes()).digest())

    def test_pdf_zielschutz_blockiert_befuellte_und_signierte_formulare(self):
        vorlage = ed.BELEHRUNG_ZIEL.with_suffix('.pdf')
        ed.prüfe_pdf_zielschutz(vorlage)
        (ROOT / '.arbeitsdaten').mkdir(exist_ok=True)
        with TemporaryDirectory(dir=ROOT / '.arbeitsdaten') as temp:
            temp = Path(temp)
            befüllt = temp / 'befüllt.pdf'
            writer = PdfWriter(clone_from=vorlage)
            writer.update_page_form_field_values(None, {'name': 'Prüfperson'}, auto_regenerate=False)
            writer.write(befüllt)
            vorher = hashlib.sha256(befüllt.read_bytes()).digest()
            with self.assertRaisesRegex(ValueError, 'befülltes oder signiertes Formularfeld'):
                ed.prüfe_pdf_zielschutz(befüllt)
            self.assertEqual(vorher, hashlib.sha256(befüllt.read_bytes()).digest())

            signiert = temp / 'signiert.pdf'
            writer = PdfWriter(clone_from=vorlage)
            signatur = next(
                ref.get_object() for seite in writer.pages for ref in seite.get('/Annots', [])
                if ref.get_object().get('/FT') == '/Sig'
            )
            signatur[NameObject('/V')] = DictionaryObject({NameObject('/Type'): NameObject('/Sig')})
            writer.write(signiert)
            with self.assertRaisesRegex(ValueError, 'befülltes oder signiertes Formularfeld'):
                ed.prüfe_pdf_zielschutz(signiert)

    def test_word_export_schuetzt_befuelltes_ziel_vor_com(self):
        skript = (ROOT / 'validierung' / 'aktualisiere_word_felder.ps1').read_text(encoding='utf-8')
        self.assertLess(skript.index('--pruefe-pdf-zielschutz'), skript.index('New-Object -ComObject Word.Application'))
        self.assertIn('ExportAsFixedFormat($pdfTemp', skript)
        self.assertLess(skript.rindex('--pruefe-pdf-zielschutz'), skript.index('Move-Item -LiteralPath $pdfTemp -Destination $pdf -Force'))
        powershell = shutil.which('pwsh') or shutil.which('powershell')
        if not powershell:
            self.skipTest('PowerShell ist nicht verfügbar.')
        (ROOT / '.arbeitsdaten').mkdir(exist_ok=True)
        with TemporaryDirectory(dir=ROOT / '.arbeitsdaten') as temp:
            pdf = Path(temp) / 'befüllt.pdf'
            writer = PdfWriter(clone_from=ed.BELEHRUNG_ZIEL.with_suffix('.pdf'))
            writer.update_page_form_field_values(None, {'name': 'Prüfperson'}, auto_regenerate=False)
            writer.write(pdf)
            docx_hash = hashlib.sha256(ed.BELEHRUNG_ZIEL.read_bytes()).digest()
            pdf_hash = hashlib.sha256(pdf.read_bytes()).digest()
            ergebnis = subprocess.run(
                [powershell, '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(ROOT / 'validierung' / 'aktualisiere_word_felder.ps1'),
                 '-DocxPfad', str(ed.BELEHRUNG_ZIEL), '-PdfPfad', str(pdf), '-PythonPfad', sys.executable,
                 '-OhneInhaltsverzeichnis'],
                capture_output=True, text=True, encoding='utf-8', errors='replace', check=False,
            )
            self.assertNotEqual(0, ergebnis.returncode)
            self.assertIn('PDF-Export abgebrochen', ergebnis.stdout + ergebnis.stderr)
            self.assertEqual(docx_hash, hashlib.sha256(ed.BELEHRUNG_ZIEL.read_bytes()).digest())
            self.assertEqual(pdf_hash, hashlib.sha256(pdf.read_bytes()).digest())


if __name__ == '__main__':
    unittest.main()
