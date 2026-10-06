#!/usr/bin/env python3
"""Validiert Status, Quellen, OSCAL-Katalog, Dokumente und Veröffentlichungsgrenzen."""

from __future__ import annotations

import argparse
import copy
import gzip
import hashlib
import io
import json
import re
import subprocess
import sys
import tarfile
import time
import unicodedata
import urllib.error
import urllib.request
import uuid
import zipfile
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import unquote
from xml.etree import ElementTree

import regex
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from jsonschema import Draft7Validator, Draft202012Validator, FormatChecker, ValidationError, validators
from pypdf import PdfReader, __version__ as PYPDF_VERSION
from erzeuge_diagramme import HÄUFIGKEITEN, RISIKOMATRIX, risikomatrix_quelle


WURZEL = Path(__file__).resolve().parents[1]
STATUS_PFAD = WURZEL / "projektstatus.json"
REGISTER_PFAD = WURZEL / "quellen" / "quellenregister.json"
REGISTER_SCHEMA_PFAD = WURZEL / "schemata" / "quellenregister.schema.json"
KATALOG_PFAD = WURZEL / "katalog" / "ki-it-sicherheitskatalog.oscal.json"
INDEX_PFAD = WURZEL / "dokumentation" / "INHALTSINDEX.md"
OSCAL_SCHEMA_PFAD = WURZEL / "schemata" / "oscal-1.1.3" / "oscal_catalog_schema.json"
DOCX_PFAD = WURZEL / "konzept" / "ki-it-sicherheitskonzept.docx"
PDF_PFAD = WURZEL / "konzept" / "ki-it-sicherheitskonzept.pdf"
BELEHRUNG_DOCX_PFAD = WURZEL / "konzept" / "anlage-1-nutzerbelehrung.docx"
BELEHRUNG_PDF_PFAD = WURZEL / "konzept" / "anlage-1-nutzerbelehrung.pdf"
BELEHRUNG_FORMULARFELDER = {"name", "organisationseinheit", "ort_datum", "unterschrift"}
DIAGRAMMQUELLEN_PFAD = WURZEL / "diagramme"
DIAGRAMMMANIFEST_PFAD = DIAGRAMMQUELLEN_PFAD / "diagramm-manifest.json"
DIAGRAMMAUSGABE_PFAD = WURZEL / "dokumentation" / "medien"

ÖFFENTLICH = "ÖFFENTLICH – organisationsneutrale Referenzvorlage"
NICHT_ÖFFENTLICH = "NICHT ÖFFENTLICH – EINSTUFUNG DURCH DIE ORGANISATION ERFORDERLICH"
PFLICHTTEILE = {"statement", "rationale", "guidance", "source"}
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{WORD_NS}}}"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
WP = f"{{{WP_NS}}}"
LAYOUT_TABELLENBREITE_DXA = 8922
LAYOUT_TABELLENEINZUG_DXA = 150
LAYOUT_ZELLENRAND_VERTIKAL_DXA = 120
LAYOUT_ZELLENRAND_HORIZONTAL_DXA = 150
DIAGRAMMSTÄMME = (
    "umgebungsuebersicht",
    "architektur",
    "architektur-cloud",
    "artefaktimport",
    "risikobewertung",
    "risikomatrix",
    "rag-datenfluss",
    "agentische-werkzeugnutzung",
)
PLANTUML_VERSION = "1.2026.6"
ABBILDUNGSBESCHRIFTUNGEN = (
    "Abbildung 1: Zwei getrennte Referenzszenarien mit gemeinsamen Nutzungsregeln.",
    "Abbildung 2: Air-Gap-Verarbeitung ohne Verbindung zu externen KI-Diensten.",
    "Abbildung 3: Cloud-Inferenz mit begrenzter Übermittlung und intern geprüften Datenrechten.",
    "Abbildung 4: Modellprüfung und Freigabe entsprechend dem Betriebsweg.",
    "Abbildung 5: Szenariobezogene Bewertung und Behandlung der KI-Risiken.",
    "Abbildung 6: Gemeinsamer Bewertungsmaßstab für beide Szenarien.",
    "Abbildung 7: Die freigegebene Wissenssuche verarbeitet bewusst abgelegte Dokumente unter durchgängigen Quellrechten.",
    "Abbildung 8: Agentenaktionen beginnen mit Einzelgenehmigungen. Projektfreigaben bleiben an Rechte und Wirkung gebunden.",
)


def lade_json(pfad: Path) -> dict[str, Any]:
    with pfad.open(encoding="utf-8") as datei:
        return json.load(datei)


def eigenschaft(quelle: dict[str, Any], name: str) -> list[str]:
    return [p["value"] for p in quelle.get("props", []) if p.get("name") == name]


def katalogkontrollen(katalog: dict[str, Any]) -> Iterable[dict[str, Any]]:
    def aus_gruppe(gruppe: dict[str, Any]) -> Iterable[dict[str, Any]]:
        yield from gruppe.get("controls", [])
        for untergruppe in gruppe.get("groups", []):
            yield from aus_gruppe(untergruppe)

    for gruppe in katalog["catalog"].get("groups", []):
        yield from aus_gruppe(gruppe)
    yield from katalog["catalog"].get("controls", [])


def oscal_validator(schema: dict[str, Any]):
    """OSCAL nutzt ECMA-262-Unicodeklassen, die Python-re nicht versteht."""

    def prüfe_muster(validator, muster, instanz, teilschema):
        if validator.is_type(instanz, "string") and regex.search(muster, instanz) is None:
            yield ValidationError(f"{instanz!r} entspricht nicht {muster!r}")

    validator_klasse = validators.extend(Draft7Validator, {"pattern": prüfe_muster})
    return validator_klasse(schema, format_checker=FormatChecker())


def prüfe_schemata(
    katalog: dict[str, Any], register: dict[str, Any]
) -> tuple[list[str], list[str]]:
    fehler: list[str] = []
    warnungen: list[str] = []
    register_schema = lade_json(REGISTER_SCHEMA_PFAD)
    for problem in sorted(
        Draft202012Validator(register_schema, format_checker=FormatChecker()).iter_errors(register),
        key=lambda eintrag: list(eintrag.path),
    ):
        ort = "/".join(str(teil) for teil in problem.path) or "Wurzel"
        fehler.append(f"Quellenregister-Schema {ort}: {problem.message}")
    oscal_schema = lade_json(OSCAL_SCHEMA_PFAD)
    for problem in sorted(oscal_validator(oscal_schema).iter_errors(katalog), key=lambda eintrag: list(eintrag.path)):
        ort = "/".join(str(teil) for teil in problem.path) or "Wurzel"
        fehler.append(f"OSCAL-1.1.3-Schema {ort}: {problem.message}")
    schema_hash = hashlib.sha256(OSCAL_SCHEMA_PFAD.read_bytes()).hexdigest()
    if schema_hash != "5e120afbd14c480a9498ab6388857ef32b3b880e458525e966ff7c7f59333d90":
        fehler.append("Das eingebundene offizielle OSCAL-1.1.3-Katalogschema besitzt eine unerwartete Prüfsumme.")
    return fehler, warnungen


def wirksame_kennzeichnung(status: dict[str, Any]) -> str:
    return NICHT_ÖFFENTLICH if status.get("organisationsspezifisch") else ÖFFENTLICH


def prüfe_projektstatus(
    status: dict[str, Any], katalog: dict[str, Any] | None = None, *, repository_modus: bool = True
) -> list[str]:
    fehler: list[str] = []
    erwartet = {
        "dokumentstatus",
        "organisationsspezifisch",
        "betriebsmodell",
        "externe-inferenz",
        "modelltraining",
        "feinabstimmung",
    }
    if set(status) != erwartet:
        fehler.append("projektstatus.json muss exakt die sechs festgelegten Steuerfelder enthalten.")
    if status.get("betriebsmodell") != "unternehmensintegriert":
        fehler.append("Das Betriebsmodell muss unternehmensintegriert sein.")
    if status.get("modelltraining") is not False:
        fehler.append("Modelltraining ist innerhalb dieses Projekts unzulässig.")
    if status.get("feinabstimmung") is not False:
        fehler.append("Feinabstimmung ist innerhalb dieses Projekts unzulässig.")
    if status.get("organisationsspezifisch") is False and status.get("dokumentstatus") != "ÖFFENTLICH":
        fehler.append("Eine organisationsneutrale Fassung muss den Dokumentstatus ÖFFENTLICH tragen.")
    if repository_modus and status.get("organisationsspezifisch") is not False:
        fehler.append("Das öffentliche Referenzrepository darf keine organisationsspezifische Fassung enthalten.")
    if repository_modus and status.get("dokumentstatus") != "ÖFFENTLICH":
        fehler.append("Das öffentliche Referenzrepository erlaubt ausschließlich den Dokumentstatus ÖFFENTLICH.")
    if repository_modus and status.get("externe-inferenz") is not False:
        fehler.append("Die öffentliche Referenz beschreibt Cloud-Nutzung, aktiviert jedoch keine externe Inferenz.")
    if status.get("externe-inferenz") and katalog:
        externe = [c for c in katalogkontrollen(katalog) if c["id"].startswith("ki-ext-")]
        if not externe or any("nicht-anwendbar" in eigenschaft(c, "standardstatus") for c in externe):
            fehler.append("Externe Inferenz ist aktiviert, während zugehörige Kontrollen noch nicht anwendbar geschaltet sind.")
    return fehler


def prüfe_quellenregister(
    register: dict[str, Any], katalog: dict[str, Any], *, lokale_dateien_erforderlich: bool = True
) -> list[str]:
    fehler: list[str] = []
    quellen = register.get("quellen", [])
    ids = [q.get("id") for q in quellen]
    uuids = [q.get("oscal_uuid") for q in quellen]
    if len(ids) != len(set(ids)):
        fehler.append("Das Quellenregister enthält doppelte Quellen-IDs.")
    if len(uuids) != len(set(uuids)):
        fehler.append("Das Quellenregister enthält doppelte OSCAL-UUIDs.")
    heute = date.today()
    for q in quellen:
        qid = q.get("id", "unbekannt")
        url = q.get("öffentliche_url", "")
        if not isinstance(url, str) or not url.startswith("https://"):
            fehler.append(f"{qid}: öffentliche HTTPS-Fundstelle fehlt.")
        if not q.get("stand") or not (q.get("veröffentlichungsdatum") or "Abrufstand" in q.get("stand", "")):
            fehler.append(f"{qid}: Veröffentlichungsdatum oder ausdrücklich dokumentierter Abrufstand fehlt.")
        if not q.get("abrufdatum") or not q.get("letzte_inhaltsprüfung"):
            fehler.append(f"{qid}: Abruf- oder Inhaltsprüfdatum fehlt.")
        if not q.get("verwendete_fundstellen"):
            fehler.append(f"{qid}: konkrete Fundstelle fehlt.")
        if q.get("lokale_fassung") and not url:
            fehler.append(f"{qid}: lokale Fassung besitzt keine öffentliche Internetfundstelle.")
        if q.get("lokale_fassung"):
            lokal = WURZEL / q["lokale_fassung"]["pfad"]
            if not lokal.is_file() and lokale_dateien_erforderlich:
                fehler.append(f"{qid}: registrierte lokale Fassung fehlt.")
            elif lokal.is_file():
                digest = hashlib.sha256(lokal.read_bytes()).hexdigest()
                if digest != q["lokale_fassung"]["sha256"]:
                    fehler.append(f"{qid}: SHA-256 der lokalen Fassung weicht ab.")
                if lokal.stat().st_size != q["lokale_fassung"]["dateigröße_bytes"]:
                    fehler.append(f"{qid}: Dateigröße der lokalen Fassung weicht ab.")
        try:
            wiedervorlage = date.fromisoformat(q["wiedervorlage_am"])
            if wiedervorlage < heute:
                fehler.append(f"{qid}: Wiedervorlage ist seit {wiedervorlage.isoformat()} überfällig.")
        except (KeyError, TypeError, ValueError):
            fehler.append(f"{qid}: Wiedervorlage ist ungültig.")

    ressourcen = katalog["catalog"].get("back-matter", {}).get("resources", [])
    ressourcen_nach_uuid = {r["uuid"]: r for r in ressourcen}
    controls = list(katalogkontrollen(katalog))
    tatsächliche_ableitung = {q["id"]: set() for q in quellen}
    quelle_nach_uuid = {q["oscal_uuid"]: q for q in quellen}
    for c in controls:
        for link in c.get("links", []):
            if not str(link.get("href", "")).startswith("#"):
                fehler.append(f"{c['id']}: Quellenlink ist kein interner Back-Matter-Verweis.")
                continue
            quelle = quelle_nach_uuid.get(link["href"][1:])
            if quelle is None:
                fehler.append(f"{c['id']}: Quellenlink kann nicht zum Quellenregister aufgelöst werden.")
            else:
                tatsächliche_ableitung[quelle["id"]].add(c["id"])
    for q in quellen:
        if q["oscal_uuid"] not in ressourcen_nach_uuid:
            fehler.append(f"{q['id']}: Back-Matter-Ressource fehlt.")
        if set(q.get("abgeleitete_kontrollen", [])) != tatsächliche_ableitung[q["id"]]:
            fehler.append(f"{q['id']}: Rückverweise auf abgeleitete Kontrollen sind nicht synchron.")
    return fehler


def prüfe_katalog(katalog: dict[str, Any]) -> list[str]:
    fehler: list[str] = []
    controls = list(katalogkontrollen(katalog))
    ids = [c.get("id") for c in controls]
    if len(ids) != len(set(ids)):
        fehler.append("Der Katalog enthält doppelte Kontroll-IDs.")
    if len(controls) < 30:
        fehler.append("Der fachliche Katalog ist für den festgelegten Erstentwurf unvollständig.")
    sichtbare: list[str] = []
    for c in controls:
        cid = c.get("id", "unbekannt")
        if not re.fullmatch(r"ki-[a-z]{3}-[0-9]{3}", cid):
            fehler.append(f"{cid}: maschinenlesbare ID entspricht nicht dem festgelegten Muster.")
        alt = eigenschaft(c, "alt-identifier")
        if alt != [cid.upper()]:
            fehler.append(f"{cid}: sichtbare Kennung fehlt oder ist unstabil.")
        sichtbare.extend(alt)
        teile = {p.get("name"): p.get("prose", "") for p in c.get("parts", [])}
        fehlend = PFLICHTTEILE - set(teile)
        if {'evidence', 'assessment-objective', 'evidence-register'} & set(teile):
            fehler.append(f"{cid}: zusätzliche Nachweisblöcke gehören nicht in das Fachkonzept.")
        szenarien = set(eigenschaft(c, 'szenario'))
        erwartet = {'cloud'} if cid == 'ki-ext-001' else {'air-gap', 'cloud'}
        if szenarien != erwartet:
            fehler.append(f"{cid}: Szenariozuordnung fehlt oder ist widersprüchlich.")
        if fehlend:
            fehler.append(f"{cid}: Pflichtteile fehlen: {', '.join(sorted(fehlend))}.")
        for abschnitt in c.get("parts", []):
            if abschnitt.get("id") != f"{cid}-{abschnitt.get('name')}":
                fehler.append(f"{cid}: stabile Abschnitts-ID fehlt oder weicht vom Kontrollnamen ab.")
        abschnitts_ids = [p.get("id") for p in c.get("parts", [])]
        if len(abschnitts_ids) != len(set(abschnitts_ids)):
            fehler.append(f"{cid}: doppelte Abschnitts-ID.")
        for name in PFLICHTTEILE & set(teile):
            if len(teile[name].strip()) < 20:
                fehler.append(f"{cid}: Teil {name} ist nicht aussagekräftig.")
        if not c.get("links"):
            fehler.append(f"{cid}: öffentliche Quellenverknüpfung fehlt.")
        if eigenschaft(c, "modalitaet") != ["MUSS"]:
            fehler.append(f"{cid}: normative Modalität MUSS fehlt.")
        if not eigenschaft(c, "verschaerfungsbegruendung"):
            fehler.append(f"{cid}: begründete projektinterne Verschärfung fehlt.")
        for pflichtprop in ("lebenszyklus", "rolle", "anwendbarkeit", "schutzbedarf", "quellenmodalitaet"):
            if not eigenschaft(c, pflichtprop):
                fehler.append(f"{cid}: kontrollierte Eigenschaft {pflichtprop} fehlt.")
        if not eigenschaft(c, "mapping-bsi-it-grundschutz") or not eigenschaft(c, "mapping-iso"):
            fehler.append(f"{cid}: orientierende BSI- oder ISO-Zuordnung fehlt.")
    if len(sichtbare) != len(set(sichtbare)):
        fehler.append("Der Katalog enthält doppelte sichtbare Kennungen.")

    nach_id = {c["id"]: c for c in controls}
    training = nach_id.get("ki-trn-001", {})
    trainingstext = " ".join(p.get("prose", "") for p in training.get("parts", []))
    if not all(begriff in trainingstext for begriff in ("Fine-Tuning", "LoRA", "RLHF", "Modellgewichte")):
        fehler.append("Das Trainings- und Feinabstimmungsverbot ist unvollständig.")
    rag_acl = " ".join(p.get("prose", "") for p in nach_id.get("ki-rag-002", {}).get("parts", []))
    rag_löschung = " ".join(p.get("prose", "") for p in nach_id.get("ki-rag-003", {}).get("parts", []))
    if "Berechtig" not in rag_acl or "Sprachmodell" not in rag_acl:
        fehler.append("Die RAG-Kontrolle enthält keine durchgängige Berechtigungsgrenze.")
    if "Löschung" not in rag_löschung or "Suchvektoren" not in rag_löschung or "Zwischenspeicher" not in rag_löschung:
        fehler.append("Die RAG-Kontrolle enthält keine vollständige Löschanforderung.")
    tooltext = " ".join(p.get("prose", "") for p in nach_id.get("ki-tol-001", {}).get("parts", []))
    if not all(begriff in tooltext for begriff in ("Lesen", "Schreiben", "Befehlsausführung", "Netzwerkzugriff")):
        fehler.append("Die Agentenkontrolle trennt Tool- und Befehlsfähigkeiten nicht vollständig.")
    externe = [c for c in controls if c["id"].startswith("ki-ext-")]
    if len(externe) != 2 or eigenschaft(nach_id.get('ki-ext-001', {}), 'anwendbarkeit') != ['Cloud-Szenario']:
        fehler.append("Die bedingte Kontrolle für externe Inferenz ist unvollständig oder falsch gekennzeichnet.")
    if eigenschaft(nach_id.get('ki-ext-002', {}), 'anwendbarkeit') != ['Beide Szenarien'] or eigenschaft(nach_id.get('ki-ext-002', {}), 'standardstatus') != ['anwendbar']:
        fehler.append("Der Schutz vor unzulässigen Ausweichzielen muss immer anwendbar sein.")
    metadata = katalog["catalog"]["metadata"]
    if set(eigenschaft(metadata, 'referenzszenario')) != {'air-gap', 'cloud'}:
        fehler.append('Der Katalog muss beide Referenzszenarien ausweisen.')
    if metadata.get("oscal-version") != "1.1.3":
        fehler.append("Der Katalog muss OSCAL 1.1.3 ausweisen.")
    if eigenschaft(metadata, "modelltraining") != ["ausgeschlossen"] or eigenschaft(metadata, "feinabstimmung") != ["ausgeschlossen"]:
        fehler.append("Metadaten weisen Trainings- oder Feinabstimmungsausschluss nicht aus.")
    return fehler


def prüfe_risikoregister(katalog: dict[str, Any]) -> list[str]:
    fehler = []
    kontrollen = {c['id']: c for c in katalogkontrollen(katalog)}
    register = [p for p in kontrollen.get('ki-gov-003', {}).get('parts', []) if p.get('name') == 'risk-register']
    if len(register) != 1:
        return ['Risikoregister: genau ein Register in KI-GOV-003 erforderlich.']
    risiken = register[0].get('parts', [])
    if {r.get('id') for r in risiken} != {f'r-{i:02d}' for i in range(1, 10)} or len(risiken) != 9:
        fehler.append('Risikoregister: stabile Kennungen R-01 bis R-09 müssen genau einmal vorkommen.')
    for risiko in risiken:
        rid = risiko.get('id')
        cloud = [p for p in risiko.get('parts', []) if p.get('name') == 'cloud-assessment']
        if len(cloud) != 1 or len(cloud[0].get('prose', '')) < 30:
            fehler.append(f'{rid}: Cloud-Bewertung und Restrisikobegründung fehlen.')
        if eigenschaft(risiko, 'szenario') != ['air-gap']:
            fehler.append(f'{rid}: Air-Gap-Zuordnung fehlt.')
        for bewertung in cloud:
            if eigenschaft(bewertung, 'szenario') != ['cloud']:
                fehler.append(f'{rid}: Cloud-Zuordnung fehlt.')
            for präfix in ('initial', 'residual'):
                werte = {p['name']: p['value'] for p in bewertung.get('props', [])}
                häufigkeit, schaden = werte.get(f'{präfix}-likelihood'), werte.get(f'{präfix}-impact')
                if häufigkeit not in HÄUFIGKEITEN or schaden not in RISIKOMATRIX:
                    fehler.append(f'{rid}: gültige Cloud-Häufigkeit und Schadenshöhe fehlen.')
                elif werte.get(f'{präfix}-risk') != RISIKOMATRIX[schaden][HÄUFIGKEITEN.index(häufigkeit)]:
                    fehler.append(f'{rid}: Cloud-Risikokategorie stimmt nicht mit der Matrix überein.')
        werte = {p['name']: p['value'] for p in risiko.get('props', [])}
        for präfix in ('initial', 'residual'):
            häufigkeit, schaden = werte.get(f'{präfix}-likelihood'), werte.get(f'{präfix}-impact')
            if häufigkeit not in HÄUFIGKEITEN or schaden not in RISIKOMATRIX:
                fehler.append(f'{rid}: gültige Häufigkeit und Schadenshöhe fehlen.')
            elif werte.get(f'{präfix}-risk') != RISIKOMATRIX[schaden][HÄUFIGKEITEN.index(häufigkeit)]:
                fehler.append(f'{rid}: Risikokategorie stimmt nicht mit der Matrix überein.')
        teile = {p['name']: p.get('prose', '') for p in risiko.get('parts', [])}
        if any(len(teile.get(n, '')) < 30 for n in ('treatment', 'assumptions', 'residual-reasoning')):
            fehler.append(f'{rid}: Behandlung, Annahmen oder Restrisikobegründung fehlen.')
        links = risiko.get('links', [])
        if not links or any(l.get('href', '')[1:] not in kontrollen or not l.get('href', '').startswith('#') for l in links):
            fehler.append(f'{rid}: verknüpfte Kontrolle fehlt.')
    if not fehler and (DIAGRAMMQUELLEN_PFAD / 'risikomatrix.puml').read_text(encoding='utf-8') != risikomatrix_quelle(katalog):
        fehler.append('Risikoregister: PlantUML-Matrix weicht vom Katalog ab.')
    return fehler


def prüfe_oscal_erweiterungen(katalog: dict[str, Any]) -> list[str]:
    fehler, ids = [], []
    namensraum = 'https://github.com/adrianweidig/allgemeines-ki-it-sicherheitskonzept/ns/oscal'
    def besuche(obj):
        if isinstance(obj, dict):
            if 'id' in obj: ids.append(obj['id'])
            for p in obj.get('props', []):
                if p['name'] not in {'alt-identifier', 'label', 'sort-id', 'status'} and p.get('ns') != namensraum:
                    fehler.append(f"OSCAL: eigener Namensraum fehlt für Eigenschaft {p['name']}.")
            for p in obj.get('parts', []):
                if p['name'] not in {'statement', 'guidance', 'assessment-objective'} and p.get('ns') != namensraum:
                    fehler.append(f"OSCAL: eigener Namensraum fehlt für Abschnitt {p['name']}.")
            for v in obj.values(): besuche(v)
        elif isinstance(obj, list):
            for v in obj: besuche(v)
    besuche(katalog)
    if len(ids) != len(set(ids)): fehler.append('OSCAL: doppelte dokumentweite ID.')
    return fehler


def prüfe_inhaltsindex(katalog: dict[str, Any]) -> list[str]:
    """Prüft die bewusst einfachen Markdown-Verweise mit expliziten HTML-Ankern."""
    fehler: list[str] = []
    if not INDEX_PFAD.is_file():
        return ["Index: Inhaltsindex fehlt."]
    indextext = INDEX_PFAD.read_text(encoding="utf-8")
    ankermuster = r'<a id="([a-z0-9-]+)"></a>'
    indexanker = re.findall(ankermuster, indextext)
    katalog_ids = {c["id"] for c in katalogkontrollen(katalog)}
    if {a for a in indexanker if a.startswith("ki-")} != katalog_ids:
        fehler.append("Index: Kontrollanker stimmen nicht mit dem Katalog überein.")
    indexlinks = [l.get("href") for l in katalog["catalog"]["metadata"].get("links", []) if l.get("rel") == "index"]
    if indexlinks != ["../dokumentation/INHALTSINDEX.md#kontrollen"] or "kontrollen" not in indexanker:
        fehler.append("Index: Katalogverweis auf den Kontrollindex fehlt oder ist ungültig.")
    texte = {INDEX_PFAD.resolve(): indextext}
    for adresse in re.findall(r"\[[^\]\n]+\]\(([^)\s]+)\)", indextext):
        if adresse.startswith(("https://", "http://")):
            continue
        pfad, _, anker = adresse.partition("#")
        ziel = (INDEX_PFAD.parent / unquote(pfad)).resolve() if pfad else INDEX_PFAD.resolve()
        if not ziel.is_relative_to(WURZEL.resolve()) or not ziel.is_file():
            fehler.append(f"Index: lokales Ziel fehlt oder liegt außerhalb des Projekts: {adresse}")
            continue
        if ziel.suffix.lower() == ".md":
            if ziel not in texte:
                texte[ziel] = ziel.read_text(encoding="utf-8")
            if anker and unquote(anker) not in re.findall(ankermuster, texte[ziel]):
                fehler.append(f"Index: Zielanker fehlt: {adresse}")
    dokumentierte_ids: set[str] = set()
    for pfad, text in texte.items():
        anker = re.findall(ankermuster, text)
        if len(anker) != len(set(anker)):
            fehler.append(f"Index: doppelte Anker in {pfad.name}.")
        dokumentierte_ids.update(a for a in anker if re.fullmatch(r"(?:a|p)\d{2}|r-\d{2}", a))
    referenzierte_ids = {i.lower() for i in re.findall(r"\b(?:A\d{2}|P\d{2}|R-\d{2})\b", indextext)}
    if referenzierte_ids != dokumentierte_ids:
        fehler.append("Index: Entscheidungs-, Risiko- oder Empfehlungskennungen sind nicht vollständig verknüpft.")
    return fehler


def _docx_text(pfad: Path) -> str:
    dokument = Document(pfad)
    teile: list[str] = []
    for abschnitt in dokument.sections:
        teile.extend(p.text for p in abschnitt.header.paragraphs)
        teile.extend(p.text for p in abschnitt.footer.paragraphs)
    teile.extend(p.text for p in dokument.paragraphs)
    for tabelle in dokument.tables:
        for zeile in tabelle.rows:
            teile.extend(zelle.text for zelle in zeile.cells)
    with zipfile.ZipFile(pfad) as paket:
        if "word/footnotes.xml" in paket.namelist():
            wurzel = ElementTree.fromstring(paket.read("word/footnotes.xml"))
            namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            teile.extend(knoten.text or "" for knoten in wurzel.iter(namespace + "t"))
    return "\n".join(teile)


def _pdf_text(pfad: Path) -> tuple[str, int]:
    leser = PdfReader(pfad)
    return "\n".join(seite.extract_text() or "" for seite in leser.pages), len(leser.pages)


def prüfe_docx_bearbeitbarkeit(pfad: Path) -> list[str]:
    """Weist Kommentare, Änderungsverfolgung, Schutz, Signaturen und Wasserzeichen ab."""
    fehler: list[str] = []
    if not zipfile.is_zipfile(pfad):
        return [f"Bearbeitbarkeit: {pfad.name} ist kein gültiges OOXML-Paket."]
    with zipfile.ZipFile(pfad) as paket:
        namen = {name.casefold() for name in paket.namelist()}
        if any(name.startswith("word/comments") or name == "word/people.xml" for name in namen):
            fehler.append(f"Bearbeitbarkeit: {pfad.name} enthält Kommentare oder zugehörige Personendaten.")
        if any(name.startswith("_xmlsignatures/") for name in namen):
            fehler.append(f"Bearbeitbarkeit: {pfad.name} enthält eine Paketsignatur.")
        xml = b"\n".join(
            paket.read(name) for name in paket.namelist()
            if name.casefold().startswith("word/") and name.casefold().endswith(".xml")
        ).lower()
        if any(marker in xml for marker in (b"commentrangestart", b"commentrangeend", b"commentreference")):
            fehler.append(f"Bearbeitbarkeit: {pfad.name} enthält Kommentarverweise.")
        if any(marker in xml for marker in (b"documentprotection", b"writeprotection")):
            fehler.append(f"Bearbeitbarkeit: {pfad.name} ist gegen Bearbeitung geschützt.")
        if b"trackrevisions" in xml:
            fehler.append(f"Bearbeitbarkeit: {pfad.name} aktiviert die Änderungsverfolgung.")
        if re.search(
            rb"<w:(?:ins|del|movefrom|moveto|cellins|celldel|cellmerge|rprchange|pprchange|tblprchange|trprchange|tcprchange|sectprchange|numberingchange)(?:\s|/?>)",
            xml,
        ):
            fehler.append(f"Bearbeitbarkeit: {pfad.name} enthält nicht angenommene Änderungen.")
        if any(marker in xml for marker in (b"watermark", b"v:textpath", b"w:background")):
            fehler.append(f"Bearbeitbarkeit: {pfad.name} enthält ein Wasserzeichen oder einen Dokumenthintergrund.")
    return fehler


def prüfe_pdf_bearbeitbarkeit(
    pfad: Path,
    *,
    erwartete_formularfelder: set[str] | None = None,
) -> list[str]:
    """Prüft PDF-Schutz, Signaturen, Formulare und redaktionelle Anmerkungen."""
    fehler: list[str] = []
    try:
        leser = PdfReader(pfad)
    except Exception as exc:
        return [f"Bearbeitbarkeit: {pfad.name} kann nicht geöffnet werden: {exc}"]
    if leser.is_encrypted:
        return [f"Bearbeitbarkeit: {pfad.name} ist verschlüsselt oder kennwortgeschützt."]
    wurzel = leser.trailer["/Root"].get_object()
    if "/Perms" in wurzel:
        fehler.append(f"Bearbeitbarkeit: {pfad.name} enthält Signatur- oder Berechtigungsbeschränkungen.")
    felder = leser.get_fields() or {}
    if erwartete_formularfelder is None:
        if "/AcroForm" in wurzel or felder:
            fehler.append(f"Bearbeitbarkeit: {pfad.name} darf keine Formularfelder enthalten.")
    elif set(felder) != erwartete_formularfelder:
        fehler.append(f"Bearbeitbarkeit: {pfad.name} enthält nicht die erwarteten Formularfelder.")
    for name, feld in felder.items():
        wert = feld.get("/V")
        if wert not in (None, "", "/Off"):
            fehler.append(f"Bearbeitbarkeit: Formularfeld {name} in {pfad.name} ist bereits befüllt oder signiert.")
    erlaubte_annotationen = {"/Link"}
    if erwartete_formularfelder is not None:
        erlaubte_annotationen.add("/Widget")
    for seite in leser.pages:
        for referenz in seite.get("/Annots", []):
            annotation = referenz.get_object()
            art = str(annotation.get("/Subtype", ""))
            if art not in erlaubte_annotationen:
                fehler.append(f"Bearbeitbarkeit: {pfad.name} enthält eine unzulässige PDF-Anmerkung ({art or 'ohne Typ'}).")
    return fehler


def prüfe_aufzählungsinterpunktion(texte: Iterable[str]) -> list[str]:
    fehler = []
    for index, text in enumerate(texte, start=1):
        bereinigt = text.rstrip()
        if bereinigt.endswith(";"):
            fehler.append(f"Aufzählungspunkt {index} endet mit einem Strichpunkt.")
    return fehler


def _xml_wert(element, name="val") -> str | None:
    if element is None:
        return None
    return element.get(W + name)


def _xml_bool_aktiv(element) -> bool:
    """Wertet ein OOXML-OnOff-Element einschließlich der Kurzform ohne Wert aus."""
    if element is None:
        return False
    return (_xml_wert(element) or "true").casefold() not in {"false", "0", "off"}


def _stil_nach_name(styles, name: str):
    """Findet eine Formatvorlage robust gegenüber lokalisierten Word-Stil-IDs."""
    for stil in styles.findall(W + "style"):
        namensfeld = stil.find(W + "name")
        if (_xml_wert(namensfeld) or "").casefold() == name.casefold():
            return stil
    return None


def _effektiver_stilwert(styles, stil, bereich: str, element: str, attribut: str) -> str | None:
    """Löst einen Stilwert über basedOn bis zu den Dokumentstandardwerten auf."""
    besucht: set[str] = set()
    aktuell = stil
    while aktuell is not None:
        stil_id = aktuell.get(W + "styleId") or ""
        if stil_id in besucht:
            break
        besucht.add(stil_id)
        knoten = aktuell.find(f"{W}{bereich}/{W}{element}")
        wert = _xml_wert(knoten, attribut)
        if wert is not None:
            return wert
        basis = aktuell.find(W + "basedOn")
        basis_id = _xml_wert(basis)
        aktuell = styles.find(f"{W}style[@{W}styleId='{basis_id}']") if basis_id else None

    standard = styles.find(f"{W}docDefaults/{W}{bereich}Default/{W}{bereich}/{W}{element}")
    return _xml_wert(standard, attribut)


def prüfe_docx_layout(pfad: Path) -> list[str]:
    """Prüft die verbindlichen maschinenlesbaren Layoutregeln des Masterdokuments."""
    fehler: list[str] = []
    dokument = Document(pfad)

    try:
        fließtext = dokument.styles["Fließtext"]
    except KeyError:
        fehler.append("Layout: die verbindliche Formatvorlage Fließtext fehlt.")
    else:
        if fließtext.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY:
            fehler.append("Layout: Fließtext muss Blocksatz verwenden.")

    fließtext_absätze = [p for p in dokument.paragraphs if p.style and p.style.name == "Fließtext"]
    if len(fließtext_absätze) < 20:
        fehler.append("Layout: zusammenhängender Konzepttext verwendet die Formatvorlage Fließtext nicht durchgängig.")

    try:
        definitionsstil = dokument.styles["Begriffsdefinition"]
    except KeyError:
        fehler.append("Layout: die verbindliche Formatvorlage Begriffsdefinition fehlt.")
    else:
        # Word entfernt eine explizite Linksbündigkeit beim Speichern, wenn sie
        # bereits dem Standard entspricht. ``None`` ist daher gleichwertig.
        if definitionsstil.paragraph_format.alignment not in (None, WD_ALIGN_PARAGRAPH.LEFT):
            fehler.append("Layout: kurze Begriffsdefinitionen müssen linksbündig gesetzt sein.")
        if definitionsstil.font.size != Pt(10.5):
            fehler.append("Layout: kurze Begriffsdefinitionen müssen Calibri 10,5 pt verwenden.")
        if definitionsstil.paragraph_format.line_spacing != 1.05:
            fehler.append("Layout: kurze Begriffsdefinitionen müssen einen Zeilenabstand von 1,05 verwenden.")
        if definitionsstil.paragraph_format.space_after != Pt(3):
            fehler.append("Layout: kurze Begriffsdefinitionen müssen 3 pt Absatzabstand danach verwenden.")
    definitionsabsätze = [p for t in dokument.tables for row in t.rows for cell in row.cells
                          for p in cell.paragraphs if p.style and p.style.name == "Begriffsdefinition"]
    if len(definitionsabsätze) < 20:
        fehler.append("Layout: das abschließende Glossar benötigt kompakte Begriffsdefinitionen.")

    if not DIAGRAMMMANIFEST_PFAD.is_file():
        fehler.append("Layout: das Prüfsummenmanifest der PlantUML-Diagramme fehlt.")
    else:
        try:
            manifest = lade_json(DIAGRAMMMANIFEST_PFAD)
        except (json.JSONDecodeError, OSError) as exc:
            fehler.append(f"Layout: das Diagrammmanifest ist nicht lesbar: {exc}")
        else:
            if manifest.get("format") != "PlantUML":
                fehler.append("Layout: das Diagrammmanifest muss PlantUML als Quellformat ausweisen.")
            if manifest.get("plantuml_version") != PLANTUML_VERSION:
                fehler.append(f"Layout: Diagramme müssen mit PlantUML {PLANTUML_VERSION} erzeugt sein.")
            if manifest.get("auflösung_png_dpi") != 180:
                fehler.append("Layout: PNG-Diagramme müssen mit 180 dpi erzeugt sein.")

            einträge = {
                eintrag.get("name"): eintrag
                for eintrag in manifest.get("diagramme", [])
                if isinstance(eintrag, dict) and eintrag.get("name")
            }
            if set(einträge) != set(DIAGRAMMSTÄMME):
                fehler.append("Layout: das Diagrammmanifest muss genau die acht Fachdiagramme enthalten.")

            for stamm in DIAGRAMMSTÄMME:
                eintrag = einträge.get(stamm, {})
                erwartete_pfade = {
                    "quelle": DIAGRAMMQUELLEN_PFAD / f"{stamm}.puml",
                    "png": DIAGRAMMAUSGABE_PFAD / f"{stamm}.png",
                    "svg": DIAGRAMMAUSGABE_PFAD / f"{stamm}.svg",
                }
                for art, diagramm_pfad in erwartete_pfade.items():
                    relativ = str(diagramm_pfad.relative_to(WURZEL)).replace("\\", "/")
                    if eintrag.get(art) != relativ:
                        fehler.append(f"Layout: Manifestpfad für {stamm}/{art} ist nicht kanonisch.")
                    if not diagramm_pfad.is_file() or diagramm_pfad.stat().st_size < 100:
                        fehler.append(f"Layout: Diagrammdatei fehlt oder ist leer: {relativ}")
                        continue
                    schlüsselfeld = f"{art}_sha256"
                    ist_hash = hashlib.sha256(diagramm_pfad.read_bytes()).hexdigest()
                    if eintrag.get(schlüsselfeld) != ist_hash:
                        fehler.append(
                            f"Layout: {relativ} stimmt nicht mit dem Diagrammmanifest überein; "
                            "Diagramme müssen vor dem Dokument neu erzeugt werden."
                        )

    abschnitt = dokument.sections[0]
    soll_cm = (21.0, 29.7, 2.5, 2.5, 2.5, 2.5, 1.25, 1.25)
    ist_cm = (
        abschnitt.page_width.cm,
        abschnitt.page_height.cm,
        abschnitt.top_margin.cm,
        abschnitt.right_margin.cm,
        abschnitt.bottom_margin.cm,
        abschnitt.left_margin.cm,
        abschnitt.header_distance.cm,
        abschnitt.footer_distance.cm,
    )
    if any(abs(soll - ist) > 0.02 for soll, ist in zip(soll_cm, ist_cm)):
        fehler.append("Layout: A4-Seite, 25-mm-Ränder oder Kopf-/Fußzeilenabstand weichen ab.")

    listen = [
        p.text for p in dokument.paragraphs
        if p.style and p.style.name in {"List Bullet", "List Number"}
    ]
    fehler.extend(prüfe_aufzählungsinterpunktion(listen))
    if any(p.text.lstrip().startswith("•") for p in dokument.paragraphs):
        fehler.append("Layout: Aufzählungen dürfen nicht als manuell gesetzte Aufzählungszeichen vorliegen.")

    with zipfile.ZipFile(pfad) as paket:
        settings = ElementTree.fromstring(paket.read("word/settings.xml"))
        auto = settings.find(W + "autoHyphenation")
        folge = settings.find(W + "consecutiveHyphenLimit")
        zone = settings.find(W + "hyphenationZone")
        if not _xml_bool_aktiv(auto):
            fehler.append("Layout: automatische Silbentrennung ist nicht aktiviert.")
        if _xml_wert(folge) != "2":
            fehler.append("Layout: höchstens zwei aufeinanderfolgende Trennzeilen müssen eingestellt sein.")
        # Word darf die explizite Trennzone beim Speichern in die gleichwertige
        # Anwendungseinstellung normalisieren. Ist sie im OOXML vorhanden, muss
        # sie weiterhin dem festgelegten Wert entsprechen.
        if zone is not None and _xml_wert(zone) != "360":
            fehler.append("Layout: die Silbentrennzone muss 360 DXA betragen.")

        styles = ElementTree.fromstring(paket.read("word/styles.xml"))
        normal_xml = _stil_nach_name(styles, "Normal")
        fließtext_xml = _stil_nach_name(styles, "Fließtext")
        if normal_xml is None:
            fehler.append("Layout: die Formatvorlage Normal fehlt.")
        else:
            schrift = _effektiver_stilwert(styles, normal_xml, "rPr", "rFonts", "ascii")
            schriftgröße = _effektiver_stilwert(styles, normal_xml, "rPr", "sz", "val")
            zeilenhöhe = _effektiver_stilwert(styles, normal_xml, "pPr", "spacing", "line")
            danach = _effektiver_stilwert(styles, normal_xml, "pPr", "spacing", "after")
            if schrift != "Calibri" or schriftgröße != "22":
                fehler.append("Layout: Normal muss Calibri 11 pt verwenden.")
            if zeilenhöhe != "276":
                fehler.append("Layout: Normal muss einen Zeilenabstand von 1,15 verwenden.")
            if danach != "120":
                fehler.append("Layout: Normal muss 6 pt Absatzabstand danach verwenden.")
        if fließtext_xml is None:
            fehler.append("Layout: die verbindliche Formatvorlage Fließtext fehlt.")
        else:
            zeilenhöhe = _effektiver_stilwert(styles, fließtext_xml, "pPr", "spacing", "line")
            danach = _effektiver_stilwert(styles, fließtext_xml, "pPr", "spacing", "after")
            if zeilenhöhe != "276":
                fehler.append("Layout: Fließtext muss einen Zeilenabstand von 1,15 verwenden.")
            if danach != "120":
                fehler.append("Layout: Fließtext muss 6 pt Absatzabstand danach verwenden.")
        sprache = _effektiver_stilwert(styles, normal_xml, "rPr", "lang", "val") if normal_xml is not None else None
        if sprache != "de-DE":
            fehler.append("Layout: die Korrektur- und Trennsprache des Grundstils muss de-DE sein.")
        for stilname in ("List Bullet", "List Number"):
            stil = _stil_nach_name(styles, stilname)
            if stil is None or stil.find(f"{W}pPr/{W}numPr") is None:
                fehler.append(f"Layout: {stilname} muss eine echte Word-Nummerierungsdefinition verwenden.")

        for stilname, größe, einzug in (("toc 1", "21", "0"), ("toc 2", "20", "340")):
            stil = _stil_nach_name(styles, stilname)
            if stil is None:
                fehler.append(f"Layout: die Inhaltsverzeichnis-Formatvorlage {stilname.upper()} fehlt.")
                continue
            schriftgröße = _effektiver_stilwert(styles, stil, "rPr", "sz", "val")
            if schriftgröße != größe:
                fehler.append(f"Layout: {stilname.upper()} besitzt nicht die festgelegte Schriftgröße.")
            absatz = stil.find(W + "pPr")
            links = absatz.find(W + "ind") if absatz is not None else None
            if (_xml_wert(links, "left") or "0") != einzug:
                fehler.append(f"Layout: {stilname.upper()} besitzt nicht den festgelegten Ebeneneinzug.")
            tabstopp = next(
                (
                    tab
                    for tab in (absatz.findall(f"{W}tabs/{W}tab") if absatz is not None else [])
                    if _xml_wert(tab) == "right"
                ),
                None,
            )
            if (
                tabstopp is None
                or abs(int(_xml_wert(tabstopp, "pos") or 0) - 9072) > 2
                or _xml_wert(tabstopp, "leader") != "dot"
            ):
                fehler.append(f"Layout: {stilname.upper()} benötigt rechtsbündige Seitenzahlen mit Punkt-Füllzeichen.")

        dokument_xml = ElementTree.fromstring(paket.read("word/document.xml"))
        feldanweisungen = [
            (element.text or "").strip()
            for element in dokument_xml.iter(W + "instrText")
            if (element.text or "").strip().upper().startswith("TOC ")
        ]
        feldanweisungen.extend(
            (element.get(W + "instr") or "").strip()
            for element in dokument_xml.iter(W + "fldSimple")
            if (element.get(W + "instr") or "").strip().upper().startswith("TOC ")
        )
        if len(feldanweisungen) != 1:
            fehler.append("Layout: das Masterdokument muss genau ein automatisch aktualisierbares Inhaltsverzeichnis enthalten.")
        elif not all(schalter in feldanweisungen[0] for schalter in ('\\o "1-2"', "\\h", "\\z", "\\u")):
            fehler.append("Layout: das Inhaltsverzeichnis muss Überschriftsebenen 1 bis 2, Verknüpfungen und Seitenzahlen verwenden.")

        toc_absätze = []
        for absatz in dokument_xml.iter(W + "p"):
            stilverweis = absatz.find(f"{W}pPr/{W}pStyle")
            stil_id = _xml_wert(stilverweis)
            stil = styles.find(f"{W}style[@{W}styleId='{stil_id}']") if stil_id else None
            stilname = _xml_wert(stil.find(W + "name")) if stil is not None else ""
            if (stilname or "").casefold() in {"toc 1", "toc 2"}:
                text = "".join(k.text or "" for k in absatz.iter(W + "t")).strip()
                toc_absätze.append(text)
        if len(toc_absätze) < 14 or any(not re.search(r"\d+$", text) for text in toc_absätze):
            fehler.append("Layout: das gespeicherte Inhaltsverzeichnis ist nicht vollständig mit rechts geführten Seitenzahlen aktualisiert.")
        if any(text in {"Dokumentenlenkung", "Inhaltsverzeichnis"} for text in toc_absätze):
            fehler.append("Layout: Vorspann oder Inhaltsverzeichnis dürfen sich nicht selbst im Inhaltsverzeichnis aufführen.")

        inline_abbildungen = dokument_xml.findall(f".//{WP}inline")
        verankerte_abbildungen = dokument_xml.findall(f".//{WP}anchor")
        if len(inline_abbildungen) != len(ABBILDUNGSBESCHRIFTUNGEN):
            fehler.append(
                f"Layout: das Masterdokument muss genau {len(ABBILDUNGSBESCHRIFTUNGEN)} inline platzierte Fachdiagramme enthalten."
            )
        if verankerte_abbildungen:
            fehler.append("Layout: frei schwebende oder verankerte Abbildungen sind unzulässig.")
        alternativtexte = []
        for inline in inline_abbildungen:
            ausdehnung = inline.find(WP + 'extent')
            if ausdehnung is not None and int(ausdehnung.get('cy', '0')) > 7380000:
                fehler.append('Layout: Diagrammhöhe überschreitet 20,5 cm einschließlich Platzreserve für die Beschriftung.')
            doc_pr = inline.find(WP + "docPr")
            alternativtext = (doc_pr.get("descr") if doc_pr is not None else "") or ""
            alternativtexte.append(alternativtext.strip())
        if any(len(text) < 80 for text in alternativtexte):
            fehler.append("Layout: jede Fachabbildung benötigt einen vollständigen Alternativtext.")

        beschriftungen = {
            p.text: p.style.name if p.style else ""
            for p in dokument.paragraphs
            if p.text.startswith("Abbildung ")
        }
        for erwartet in ABBILDUNGSBESCHRIFTUNGEN:
            if beschriftungen.get(erwartet) != "Abbildungsbeschriftung":
                fehler.append(f"Layout: Abbildungsbeschriftung fehlt oder verwendet den falschen Stil: {erwartet}")

        tabellen = dokument_xml.findall(f".//{W}tbl")
        if not tabellen:
            fehler.append("Layout: erwartete Datentabellen fehlen.")
        for tabellenindex, tabelle in enumerate(tabellen, start=1):
            tbl_pr = tabelle.find(W + "tblPr")
            tbl_w = tbl_pr.find(W + "tblW") if tbl_pr is not None else None
            tbl_ind = tbl_pr.find(W + "tblInd") if tbl_pr is not None else None
            layout = tbl_pr.find(W + "tblLayout") if tbl_pr is not None else None
            ausrichtung = tbl_pr.find(W + "jc") if tbl_pr is not None else None
            if _xml_wert(tbl_w, "type") != "dxa" or _xml_wert(tbl_w, "w") != str(LAYOUT_TABELLENBREITE_DXA):
                fehler.append(f"Layout: Tabelle {tabellenindex} besitzt keine feste Breite von {LAYOUT_TABELLENBREITE_DXA} DXA.")
            if _xml_wert(tbl_ind, "type") != "dxa" or _xml_wert(tbl_ind, "w") != str(LAYOUT_TABELLENEINZUG_DXA):
                fehler.append(f"Layout: Tabelle {tabellenindex} besitzt nicht den festgelegten Einzug.")
            if _xml_wert(layout, "type") != "fixed":
                fehler.append(f"Layout: Tabelle {tabellenindex} verwendet keine feste Geometrie.")
            # Fehlendes w:jc ist nach OOXML die Word-Standardeinstellung links.
            if _xml_wert(ausrichtung) not in {None, "left", "start"}:
                fehler.append(f"Layout: Tabelle {tabellenindex} ist nicht linksbündig ausgerichtet.")

            grid = tabelle.find(W + "tblGrid")
            spalten = [int(_xml_wert(e, "w") or 0) for e in (list(grid) if grid is not None else [])]
            if not spalten or sum(spalten) != LAYOUT_TABELLENBREITE_DXA:
                fehler.append(f"Layout: Spaltenraster der Tabelle {tabellenindex} ist nicht breitenkonsistent.")
            zeilen = tabelle.findall(W + "tr")
            if not zeilen or zeilen[0].find(f"{W}trPr/{W}tblHeader") is None:
                fehler.append(f"Layout: Kopfzeile der Tabelle {tabellenindex} ist nicht als Wiederholungszeile markiert.")
            if any(_xml_wert(h, "hRule") == "exact" for h in tabelle.findall(f".//{W}trHeight")):
                fehler.append(f"Layout: Tabelle {tabellenindex} verwendet eine unzulässige exakte Zeilenhöhe.")

            for zeilenindex, zeile in enumerate(zeilen, start=1):
                zellen = zeile.findall(W + "tc")
                for spaltenindex, zelle in enumerate(zellen):
                    tc_pr = zelle.find(W + "tcPr")
                    tc_w = tc_pr.find(W + "tcW") if tc_pr is not None else None
                    if spaltenindex >= len(spalten) or _xml_wert(tc_w, "w") != str(spalten[spaltenindex]):
                        fehler.append(
                            f"Layout: Zellenbreite in Tabelle {tabellenindex}, Zeile {zeilenindex}, "
                            f"Spalte {spaltenindex + 1} weicht vom Raster ab."
                        )
                    v_align = tc_pr.find(W + "vAlign") if tc_pr is not None else None
                    if _xml_wert(v_align) != "center":
                        fehler.append(
                            f"Layout: Zelle in Tabelle {tabellenindex}, Zeile {zeilenindex}, "
                            f"Spalte {spaltenindex + 1} ist vertikal nicht zentriert."
                        )
                    ränder = tc_pr.find(W + "tcMar") if tc_pr is not None else None
                    mindestwerte = (
                        (("top",), LAYOUT_ZELLENRAND_VERTIKAL_DXA),
                        (("bottom",), LAYOUT_ZELLENRAND_VERTIKAL_DXA),
                        (("start", "left"), LAYOUT_ZELLENRAND_HORIZONTAL_DXA),
                        (("end", "right"), LAYOUT_ZELLENRAND_HORIZONTAL_DXA),
                    )
                    for namen, minimum in mindestwerte:
                        rand = next(
                            (ränder.find(W + name) for name in namen if ränder is not None and ränder.find(W + name) is not None),
                            None,
                        )
                        if _xml_wert(rand, "type") != "dxa" or int(_xml_wert(rand, "w") or 0) < minimum:
                            fehler.append(
                                f"Layout: Zellrand {'/'.join(namen)} in Tabelle {tabellenindex}, Zeile {zeilenindex}, "
                                f"Spalte {spaltenindex + 1} ist zu klein."
                            )

        if dokument.tables:
            abkürzungstabelle = dokument.tables[-1]
            kopf = [zelle.text.strip() for zelle in abkürzungstabelle.rows[0].cells]
            if kopf != ["Begriff oder Abkürzung", "Bedeutung im Konzept"]:
                fehler.append(
                    "Layout: das Glossar muss Begriffe und Erläuterungen in zwei Spalten enthalten."
                )
            if len(abkürzungstabelle.rows) < 20:
                fehler.append(
                    "Layout: das Glossar deckt die verwendeten Fachbegriffe nicht ausreichend ab."
                )

        footer_namen = [name for name in paket.namelist() if re.fullmatch(r"word/footer\d+\.xml", name)]
        footer_text = ""
        footer_felder: list[str] = []
        rechtsstopp = False
        obere_linie = False
        for name in footer_namen:
            wurzel = ElementTree.fromstring(paket.read(name))
            footer_text += " ".join(k.text or "" for k in wurzel.iter(W + "t"))
            footer_felder.extend((k.get(W + "instr") or "").strip() for k in wurzel.iter(W + "fldSimple"))
            footer_felder.extend((k.text or "").strip() for k in wurzel.iter(W + "instrText"))
            rechtsstopp = rechtsstopp or any(
                _xml_wert(k) == "right"
                and abs(int(_xml_wert(k, "pos") or 0) - 9072) <= 2
                for k in wurzel.iter(W + "tab")
            )
            obere_linie = obere_linie or any(_xml_wert(k) == "single" for k in wurzel.iter(W + "top"))
        if not {"PAGE", "NUMPAGES"}.issubset(set(footer_felder)):
            fehler.append("Layout: Fußzeile muss PAGE und NUMPAGES als Felder enthalten.")
        if "organisationsneutrale Referenzvorlage" in footer_text:
            fehler.append("Layout: die lange Schutzkennzeichnung darf die Fußzeile nicht überladen.")
        if not rechtsstopp or not obere_linie:
            fehler.append("Layout: rechte Seitenführung oder dezente obere Fußzeilenlinie fehlt.")

    return fehler


def prüfe_pdf_seitenführung(pfad: Path) -> list[str]:
    fehler: list[str] = []
    leser = PdfReader(pfad)
    seitenzahl = len(leser.pages)
    for nummer, seite in enumerate(leser.pages, start=1):
        if nummer == 1:
            continue
        text = re.sub(r"\s+", " ", seite.extract_text() or "")
        if f"Seite {nummer} von {seitenzahl}" not in text:
            fehler.append(f"Layout: PDF-Seite {nummer} besitzt keine korrekte Seitenführung.")
    return fehler


def normalisiere_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).replace("\u00ad", "")
    # Nur durch den Seitenumbruch erzeugte Trennstellen zusammenführen. Ein
    # absichtlich gesetzter Bindestrich mit folgendem Leerzeichen, etwa in
    # „Modell- und Laufzeit“, muss für den Inhaltsvergleich erhalten bleiben.
    text = re.sub(r"(?<=\w)-[ \t]*\r?\n[ \t]*(?=\w)", "", text)
    text = re.sub(r"\bSeite\s+\d+\b", " ", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip().casefold()


def prüfe_versionsgleichheit(katalog_version: str, docx_text: str, pdf_text: str) -> list[str]:
    fehler = []
    marker = f"Version {katalog_version}"
    if marker not in docx_text:
        fehler.append("Die DOCX-Fassung weist nicht dieselbe Version wie der Katalog aus.")
    if marker not in pdf_text:
        fehler.append("Die PDF-Fassung weist nicht dieselbe Version wie der Katalog aus.")
    return fehler


def prüfe_konzepttrennung(text: str) -> list[str]:
    """Hält Übernahmehinweise, Redaktionssprache und unnötigen Fachjargon aus dem Fachkonzept heraus."""
    fehler: list[str] = []
    verbotene_passagen = (
        "Verfahren zur organisationsspezifischen Übernahme",
        "Referenzrepository",
        "vor der ersten Organisationsangabe",
        "übernehmenden Organisation",
        "Die Matrix ist ein Prüfungseinstieg und keine Rechtsberatung",
        "Die folgenden Kontrollblöcke übernehmen",
        "Der vollständige Wortlaut, die Quellenlinks und sämtliche Eigenschaften verbleiben",
        "Maßgeblich ist das maschinenlesbare Quellenregister",
        "Inhaltsverzeichnis wird beim Erzeugen der Lesefassung aktualisiert",
        "Keine amtliche Veröffentlichung · keine Systemfreigabe · kein Zertifikat",
        "Das Dokument ist wie folgt aufgebaut",
        "Dieses Kapitel beschreibt den Aufbau",
        "Die folgende Tabelle zeigt",
        "Die nachfolgende Tabelle zeigt",
        "Der Katalog enthält die folgenden",
    )
    for passage in verbotene_passagen:
        if passage.casefold() in text.casefold():
            fehler.append(f"Konzepttrennung: Übernahmeanweisung im Fachkonzept gefunden: {passage}")
    unerklärter_fachjargon = (
        "Provenienz",
        "Poisoning",
        "untrusted",
        "Least Privilege",
        "Tool-Policy",
        "Circuit Breaker",
        "Queueing",
        "Exposure",
        "Baseline",
        "Rollback",
        "Regression",
        "Reranking",
        "Drift",
        "KI-Governance",
        "Reviewprotokolle",
        "Reviews",
        "Geheimschutzgate",
        "Geheimschutzgates",
        "On-Demand-Uploads",
        "Tool-Fähigkeiten",
        "Tool-Aufruf",
        "Tool-Nutzung",
        "Tool-Protokollierung",
        "Toolgrenzen",
        "Toolkette",
        "Tool-Verantwortliche",
        "Auditdaten",
        "Auditnachweise",
        "Artefakthashes",
        "Mappingübersicht",
        "automatischer externer Fallback",
        "Cloud-Fallbacks",
        "Provider-Erkennung",
        "Konfigurationsdrift",
        "Hostalarmen",
        "Parser",
        "Caches",
    )
    for begriff in unerklärter_fachjargon:
        if re.search(rf"\b{re.escape(begriff)}\b", text, re.IGNORECASE):
            fehler.append(f"Konzepttrennung: vermeidbarer oder nicht vorab erklärter Fachbegriff gefunden: {begriff}")
    if "14.1 Quellenverzeichnis" not in text or "14.2 Glossar und Abkürzungen" not in text:
        fehler.append("Konzepttrennung: Kapitel 14 muss unmittelbar Quellenverzeichnis und Abkürzungen enthalten.")
    if "2.1 Zentrale Begriffe" in text:
        fehler.append("Konzepttrennung: Begriffserklärungen gehören in das Glossar am Ende.")
    for risiko_id in (f"R-{nummer:02d}" for nummer in range(1, 10)):
        if risiko_id not in text:
            fehler.append(f"Konzepttrennung: Risikoszenario {risiko_id} fehlt im Risikoregister.")
    if re.search(r"\b15(?:\.|\s)\s*(?:Quellenverzeichnis|Glossar|Abkürzungen)", text):
        fehler.append("Konzepttrennung: ein veraltetes Kapitel 15 ist im Fachkonzept verblieben.")
    return fehler


def prüfe_dokumente(katalog: dict[str, Any], status: dict[str, Any]) -> tuple[list[str], list[str]]:
    fehler: list[str] = []
    warnungen: list[str] = []
    for pfad in (DOCX_PFAD, PDF_PFAD, BELEHRUNG_DOCX_PFAD, BELEHRUNG_PDF_PFAD):
        if not pfad.is_file() or pfad.stat().st_size < 1000:
            fehler.append(f"Dokument fehlt oder ist technisch leer: {pfad.relative_to(WURZEL)}")
    if fehler:
        return fehler, warnungen
    if not zipfile.is_zipfile(DOCX_PFAD):
        fehler.append("Das DOCX ist kein gültiges OOXML-ZIP-Paket.")
        return fehler, warnungen
    try:
        docx_text = _docx_text(DOCX_PFAD)
    except Exception as exc:
        fehler.append(f"DOCX kann nicht technisch geöffnet werden: {exc}")
        return fehler, warnungen
    try:
        pdf_text, seiten = _pdf_text(PDF_PFAD)
    except Exception as exc:
        fehler.append(f"PDF kann nicht technisch geöffnet werden: {exc}")
        return fehler, warnungen
    if seiten < 10:
        fehler.append("Die PDF-Fassung ist für den festgelegten Fachinhalt unerwartet kurz.")
    fehler.extend(prüfe_docx_bearbeitbarkeit(DOCX_PFAD))
    fehler.extend(prüfe_docx_bearbeitbarkeit(BELEHRUNG_DOCX_PFAD))
    fehler.extend(prüfe_pdf_bearbeitbarkeit(PDF_PFAD))
    fehler.extend(prüfe_pdf_bearbeitbarkeit(
        BELEHRUNG_PDF_PFAD,
        erwartete_formularfelder=BELEHRUNG_FORMULARFELDER,
    ))
    fehler.extend(prüfe_docx_layout(DOCX_PFAD))
    fehler.extend(prüfe_pdf_seitenführung(PDF_PFAD))
    version = katalog["catalog"]["metadata"]["version"]
    fehler.extend(prüfe_versionsgleichheit(version, docx_text, pdf_text))
    fehler.extend(prüfe_konzepttrennung(docx_text))
    kennzeichnung = wirksame_kennzeichnung(status)
    if kennzeichnung not in docx_text or kennzeichnung not in pdf_text:
        fehler.append("Schutzkennzeichnung ist in DOCX und PDF nicht konsistent wirksam.")
    for c in katalogkontrollen(katalog):
        kennung = c["id"].upper()
        if kennung not in docx_text:
            fehler.append(f"DOCX enthält die Kontroll-ID {kennung} nicht.")
        if kennung not in pdf_text:
            fehler.append(f"PDF enthält die Kontroll-ID {kennung} nicht.")
    norm_docx = normalisiere_text(docx_text)
    norm_pdf = normalisiere_text(pdf_text)
    fehler.extend(prüfe_katalogableitung(katalog, docx_text))
    # Der Prosaabgleich bewertet alphabetische Wörter. Veränderliche URLs,
    # Hashes, Datumswerte und technische IDs werden separat validiert und
    # würden wegen unterschiedlicher PDF-Zeilenumbrüche nur Scheindifferenzen
    # erzeugen.
    wortmuster = r"(?<![\w-])[a-zäöüß]+(?![\w-])"
    docx_token = re.findall(wortmuster, norm_docx)
    pdf_token = re.findall(wortmuster, norm_pdf)
    docx_häufigkeit = Counter(docx_token)
    pdf_häufigkeit = Counter(pdf_token)
    gemeinsame_token = sum((docx_häufigkeit & pdf_häufigkeit).values())
    tokenabdeckung = gemeinsame_token / max(1, len(docx_token))
    docx_wörter = {wort for wort in docx_häufigkeit if len(wort) >= 5}
    pdf_wörter = {wort for wort in pdf_häufigkeit if len(wort) >= 5}
    wortabdeckung = len(docx_wörter & pdf_wörter) / max(1, len(docx_wörter))
    # Fußnoten stehen im OOXML technisch gesammelt am Dokumentende, im PDF aber
    # auf der jeweiligen Seite. Ein globaler Sequenzvergleich wäre deshalb
    # fachlich ungeeignet und bei langen Dokumenten quadratisch langsam.
    if tokenabdeckung < 0.98 or wortabdeckung < 0.96:
        fehler.append(
            "Normalisierter DOCX/PDF-Text weicht zu stark ab "
            f"(Tokenabdeckung {tokenabdeckung:.3f}, Wortabdeckung {wortabdeckung:.3f})."
        )
    return fehler, warnungen


def prüfe_katalogableitung(katalog: dict[str, Any], docx_text: str) -> list[str]:
    """Jeder fachliche Kontroll- und Risikoabschnitt muss im Master enthalten sein."""
    fehler = []
    text = normalisiere_text(docx_text)
    def prüfe_teile(teile):
        for teil in teile:
            if teil['name'] not in {'source', 'rationale'} and normalisiere_text(teil.get('prose', '')) not in text:
                fehler.append(f"Katalogableitung: Abschnitt {teil['id']} fehlt im DOCX oder weicht ab.")
            prüfe_teile(teil.get('parts', []))
    for kontrolle in katalogkontrollen(katalog): prüfe_teile(kontrolle.get('parts', []))
    return fehler


def _menschenlesbare_jsontexte(objekt: Any, pfad: tuple[Any, ...] = ()) -> Iterable[str]:
    if isinstance(objekt, dict):
        for schlüssel, wert in objekt.items():
            if schlüssel in {"öffentliche_url", "href", "$schema", "$id", "pfad", "sha256", "offizielle_sha256_am_prüftag", "normalisierter_text_sha256_beider_fassungen"}:
                continue
            if schlüssel == "value" and "props" in pfad:
                continue
            yield from _menschenlesbare_jsontexte(wert, pfad + (schlüssel,))
    elif isinstance(objekt, list):
        for index, wert in enumerate(objekt):
            yield from _menschenlesbare_jsontexte(wert, pfad + (index,))
    elif isinstance(objekt, str):
        yield objekt


def prüfe_inhaltsgrenzen(katalog: dict[str, Any], register: dict[str, Any]) -> tuple[list[str], list[str]]:
    fehler: list[str] = []
    warnungen: list[str] = []
    texte: list[tuple[str, str]] = []
    for pfad in [
        WURZEL / "README.md",
        WURZEL / "AGENTS.md",
        WURZEL / "SECURITY.md",
        WURZEL / "SUPPORT.md",
        WURZEL / "CONTRIBUTING.md",
        WURZEL / "CODE_OF_CONDUCT.md",
        WURZEL / "CHANGELOG.md",
        *sorted((WURZEL / "dokumentation").rglob("*.md")),
    ]:
        if pfad.is_file():
            texte.append((str(pfad.relative_to(WURZEL)), pfad.read_text(encoding="utf-8")))
    texte.append(("OSCAL-Katalog", "\n".join(_menschenlesbare_jsontexte(katalog))))
    texte.append(("Quellenregister", "\n".join(_menschenlesbare_jsontexte(register))))
    unerwünscht = [r"\bfuer\b", r"\boeffentlich\b", r"\bueberpruefen\b"]
    platzhalter = [r"\bTODO\b", r"\bTBD\b", r"change-me", r"your-name", r"your-repo", r"example\.com", r"lorem ipsum", r"coming soon", r"\bplaceholder\b"]
    positive_konformität = [r"garantiert(?:e|er|es)? Konformität", r"vollständig ISO-konform", r"automatisch BSI-konform", r"dieses Projekt ist zertifiziert"]
    geheimnisse = [r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----", r"\bgh[pousr]_[A-Za-z0-9]{30,}\b", r"\bsk-[A-Za-z0-9]{24,}\b"]
    for name, text in texte:
        verbotener_status = "OF" + "FEN"
        if re.search(rf"\b{verbotener_status}\b", text):
            fehler.append(f"{name}: unzulässige Dokumentstatus-Bezeichnung gefunden.")
        for muster in unerwünscht:
            if re.search(muster, text, re.IGNORECASE):
                fehler.append(f"{name}: unerwünschte deutsche Ersatzschreibweise zu Muster {muster}.")
        for muster in platzhalter:
            if re.search(muster, text, re.IGNORECASE):
                fehler.append(f"{name}: verbotener Platzhalter zu Muster {muster}.")
        for muster in positive_konformität:
            if re.search(muster, text, re.IGNORECASE):
                fehler.append(f"{name}: unzulässige Konformitätsbehauptung.")
        for muster in geheimnisse:
            if re.search(muster, text):
                fehler.append(f"{name}: mögliches Geheimnis gefunden.")
        bereinigt = re.sub(r"https://\S+", "", text)
        if re.search(r"(?<!\d)(?:10|127|169\.254|172\.(?:1[6-9]|2\d|3[01])|192\.168)\.\d{1,3}\.\d{1,3}(?!\d)", bereinigt):
            fehler.append(f"{name}: mögliche reale interne IP-Adresse gefunden.")
        if re.search(r"\b[a-z0-9-]+\.(?:local|intern|intra|corp)\b", bereinigt, re.IGNORECASE):
            fehler.append(f"{name}: möglicher organisationsspezifischer Host- oder Domänenname gefunden.")
    return fehler, warnungen


def prüfe_veröffentlichungsliste(pfade: Iterable[str]) -> list[str]:
    fehler = []
    normiert = [p.replace("\\", "/") for p in pfade]
    if any(p.startswith("quellen/lokale-eingaben/") for p in normiert):
        fehler.append("Eine lokale Eingangsdatei ist im Git- oder Veröffentlichungsbestand enthalten.")
    if any(p.startswith(".arbeitsdaten/") for p in normiert):
        fehler.append("Temporäre Arbeitsdaten sind im Git- oder Veröffentlichungsbestand enthalten.")
    return fehler


def _git(*argumente: str, binär: bool = False) -> subprocess.CompletedProcess:
    return subprocess.run(
        ["git", *argumente], cwd=WURZEL, check=False,
        capture_output=True, text=not binär, encoding=None if binär else "utf-8",
    )


def prüfe_git() -> tuple[list[str], list[str]]:
    fehler: list[str] = []
    warnungen: list[str] = []
    if not (WURZEL / ".git").is_dir():
        warnungen.append("Git-Prüfung übersprungen: Repository ist noch nicht initialisiert.")
        return fehler, warnungen
    lokale_pdfs = sorted((WURZEL / "quellen" / "lokale-eingaben").glob("*.pdf"))
    for pfad in lokale_pdfs:
        relativ = str(pfad.relative_to(WURZEL)).replace("\\", "/")
        ergebnis = _git("check-ignore", "-q", "--", relativ)
        if ergebnis.returncode != 0:
            fehler.append(f"Lokale Eingabe wird nicht von Git ignoriert: {relativ}")
    gelistet = _git("ls-files", "-z", binär=True)
    if gelistet.returncode != 0:
        fehler.append("git ls-files konnte nicht ausgeführt werden.")
    else:
        pfade = [p.decode("utf-8") for p in gelistet.stdout.split(b"\0") if p]
        fehler.extend(prüfe_veröffentlichungsliste(pfade))
    kopf = _git("rev-parse", "--verify", "HEAD")
    if kopf.returncode == 0:
        archiv = _git("archive", "--format=tar", "HEAD", binär=True)
        if archiv.returncode != 0:
            fehler.append("git archive konnte nicht erzeugt werden.")
        else:
            with tarfile.open(fileobj=io.BytesIO(archiv.stdout), mode="r:") as paket:
                fehler.extend(prüfe_veröffentlichungsliste(paket.getnames()))
    else:
        warnungen.append("Archivprüfung übersprungen: Es existiert noch kein Commit.")
    return fehler, warnungen


def _hole_url(url: str, *, vollständig: bool = False, versuche: int = 2) -> tuple[int | None, bytes, str | None]:
    letzter_fehler: str | None = None
    for versuch in range(versuche):
        try:
            kopfzeilen = {"User-Agent": "KI-Sicherheitskonzept-Quellenpruefung/0.1 (+https://github.com/adrianweidig/allgemeines-ki-it-sicherheitskonzept)"}
            if not vollständig:
                kopfzeilen["Range"] = "bytes=0-4095"
            anfrage = urllib.request.Request(url, headers=kopfzeilen, method="GET")
            with urllib.request.urlopen(anfrage, timeout=25) as antwort:
                daten = antwort.read() if vollständig else antwort.read(4096)
                if vollständig and "gzip" in (antwort.headers.get("Content-Encoding") or "").lower():
                    daten = gzip.decompress(daten)
                return antwort.status, daten, None
        except urllib.error.HTTPError as exc:
            if exc.code in {403, 405, 429}:
                return exc.code, b"", str(exc)
            letzter_fehler = str(exc)
            if exc.code in {404, 410}:
                return exc.code, b"", letzter_fehler
        except Exception as exc:
            letzter_fehler = str(exc)
        if versuch + 1 < versuche:
            time.sleep(1.0)
    return None, b"", letzter_fehler


def _normalisierter_pdf_text(daten: bytes) -> str:
    text = "\n".join(seite.extract_text() or "" for seite in PdfReader(io.BytesIO(daten)).pages)
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", text)).strip()


def prüfe_onlinequellen(register: dict[str, Any]) -> tuple[list[str], list[str]]:
    fehler: list[str] = []
    warnungen: list[str] = []
    quellen = register["quellen"]
    with ThreadPoolExecutor(max_workers=8) as executor:
        aufgaben = {executor.submit(_hole_url, q["öffentliche_url"]): q for q in quellen}
        for aufgabe in as_completed(aufgaben):
            q = aufgaben[aufgabe]
            status, _, meldung = aufgabe.result()
            if status in {404, 410}:
                fehler.append(f"{q['id']}: öffentliche Quelle antwortet mit HTTP {status}.")
            elif status in {403, 405, 429}:
                warnungen.append(f"{q['id']}: HTTP {status}; manuelle Erreichbarkeitsprüfung erforderlich.")
            elif status is None or status >= 500:
                warnungen.append(f"{q['id']}: temporärer Abruffehler; manuelle Prüfung erforderlich ({meldung}).")
            elif status < 200 or status >= 400:
                fehler.append(f"{q['id']}: unerwarteter HTTP-Status {status}.")

    for q in quellen:
        lokal = q.get("lokale_fassung")
        if not lokal:
            continue
        pdf_url = lokal.get("offizielle_pdf_url") or q["öffentliche_url"]
        status, daten, meldung = _hole_url(pdf_url, vollständig=True)
        if status is None or status >= 400:
            if status in {403, 405, 429} or status is None or (status and status >= 500):
                warnungen.append(f"{q['id']}: offizieller PDF-Abgleich temporär nicht möglich ({status or meldung}).")
                continue
            fehler.append(f"{q['id']}: offizieller PDF-Abgleich scheitert mit HTTP {status}.")
            continue
        remote_hash = hashlib.sha256(daten).hexdigest()
        erwartet = lokal.get("offizielle_sha256_am_prüftag")
        if remote_hash != erwartet:
            fehler.append(f"{q['id']}: offizielle PDF-Prüfsumme hat sich seit der Inhaltsprüfung geändert.")
            continue
        if remote_hash != lokal["sha256"]:
            text_hash = hashlib.sha256(_normalisierter_pdf_text(daten).encode("utf-8")).hexdigest()
            registriert = lokal.get("normalisierter_text_sha256_beider_fassungen")
            verfahren = lokal.get("normalisierungsverfahren", "")
            lokaler_pfad = WURZEL / lokal["pfad"]
            lokal_text_hash = registriert
            if lokaler_pfad.is_file():
                lokal_daten = lokaler_pfad.read_bytes()
                lokal_text_hash = hashlib.sha256(_normalisierter_pdf_text(lokal_daten).encode("utf-8")).hexdigest()
            if f"pypdf {PYPDF_VERSION}" not in verfahren:
                fehler.append(f"{q['id']}: Normalisierungsverfahren passt nicht zur installierten pypdf-Version.")
            elif not registriert or text_hash != lokal_text_hash or text_hash != registriert:
                fehler.append(f"{q['id']}: binär abweichende offizielle PDF ist nicht mehr textidentisch zur lokalen Fassung.")
    return fehler, warnungen


def führe_prüfungen_aus(
    *, streng: bool, online: bool, repository_modus: bool, lokale_dateien_erforderlich: bool
) -> int:
    fehler: list[str] = []
    warnungen: list[str] = []
    status = lade_json(STATUS_PFAD)
    register = lade_json(REGISTER_PFAD)
    katalog = lade_json(KATALOG_PFAD)

    for funktion, argumente in [
        (prüfe_schemata, (katalog, register)),
        (prüfe_quellenregister, (register, katalog)),
        (prüfe_katalog, (katalog,)),
        (prüfe_risikoregister, (katalog,)),
        (prüfe_oscal_erweiterungen, (katalog,)),
        (prüfe_inhaltsindex, (katalog,)),
        (prüfe_projektstatus, (status, katalog)),
        (prüfe_inhaltsgrenzen, (katalog, register)),
        (prüfe_git, ()),
    ]:
        if funktion is prüfe_projektstatus:
            ergebnis = funktion(*argumente, repository_modus=repository_modus)
        elif funktion is prüfe_quellenregister:
            ergebnis = funktion(*argumente, lokale_dateien_erforderlich=lokale_dateien_erforderlich)
        else:
            ergebnis = funktion(*argumente)
        if isinstance(ergebnis, tuple):
            neue_fehler, neue_warnungen = ergebnis
            fehler.extend(neue_fehler)
            warnungen.extend(neue_warnungen)
        else:
            fehler.extend(ergebnis)

    dokument_fehler, dokument_warnungen = prüfe_dokumente(katalog, status)
    fehler.extend(dokument_fehler)
    warnungen.extend(dokument_warnungen)
    if online:
        online_fehler, online_warnungen = prüfe_onlinequellen(register)
        fehler.extend(online_fehler)
        warnungen.extend(online_warnungen)

    print(f"Validierung: {len(fehler)} Fehler, {len(warnungen)} Warnungen")
    for meldung in warnungen:
        print(f"WARNUNG: {meldung}")
    for meldung in fehler:
        print(f"FEHLER: {meldung}")
    if fehler or (streng and any("überfällig" in w.lower() for w in warnungen)):
        return 1
    print("ERGEBNIS: alle verbindlichen Prüfungen bestanden")
    return 0


def parser() -> argparse.ArgumentParser:
    befehle = argparse.ArgumentParser(description=__doc__)
    befehle.add_argument("--streng", action="store_true", help="Verbindlicher Prüfmodus für Veröffentlichungen.")
    befehle.add_argument("--online", action="store_true", help="Öffentliche Quellen und offizielle PDF-Fassungen online prüfen.")
    befehle.add_argument(
        "--ohne-lokale-eingaben", action="store_true",
        help="Nur für CI-Checkouts: registrierte Metadaten und amtliche Online-Fassungen ohne ignorierte lokale PDFs prüfen.",
    )
    befehle.add_argument(
        "--offline-anpassung", action="store_true",
        help="Organisationsspezifischen Offline-Modus erlauben; niemals im öffentlichen Repository verwenden.",
    )
    return befehle


if __name__ == "__main__":
    argumente = parser().parse_args()
    sys.exit(führe_prüfungen_aus(
        streng=argumente.streng,
        online=argumente.online,
        repository_modus=not argumente.offline_anpassung,
        lokale_dateien_erforderlich=not argumente.ohne_lokale_eingaben,
    ))
